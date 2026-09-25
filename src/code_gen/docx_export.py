"""Geração da Ata Oficial de Sorteio em formato Word (.docx) editável."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor


def generate_draw_report_docx(
    output_path: Path | str,
    drawn_records: list[dict[str, str]],
    *,
    institution_name: str = "Colégio Adventista de Blumenau",
    event_title: str = "Sorteio da Rematrícula",
) -> Path:
    """Gera um arquivo Word (.docx) contendo a Ata Oficial do Sorteio.

    A tabela de vencedores inclui uma coluna em branco 'Prêmio Recebido' para que
    a instituição possa editar e preencher o prêmio entregue a cada ganhador.
    """
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()

    # Definir margens do documento (2 cm)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Título Principal
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title_p.add_run(f"ATA DE REALIZAÇÃO DE SORTEIO — {event_title.upper()}")
    run_title.bold = True
    run_title.font.name = "Arial"
    run_title.font.size = Pt(16)
    run_title.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    # Subtítulo Institucional
    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = sub_p.add_run(institution_name)
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(12)
    run_sub.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    doc.add_paragraph()

    # Informações da Apuração
    now_str = datetime.now().strftime("%d/%m/%Y às %H:%M:%S")
    info_p = doc.add_paragraph()
    run_info = info_p.add_run(
        f"Aos {datetime.now().strftime('%d')} dias do mês de {datetime.now().strftime('%B')} "
        f"de {datetime.now().year}, às {now_str.split(' às ')[1]}, foi realizada a apuração "
        f"oficial do {event_title}. Abaixo constam os participantes contemplados e os "
        "respectivos códigos sorteados."
    )
    run_info.font.name = "Arial"
    run_info.font.size = Pt(11)

    doc.add_paragraph()

    # Tabela de Resultados
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Cabeçalho da Tabela
    hdr_cells = table.rows[0].cells
    headers = ["№", "Aluno", "E-mail / Responsável", "Código Sorteado", "Prêmio Recebido"]
    widths = [Inches(0.5), Inches(2.2), Inches(2.2), Inches(1.3), Inches(2.2)]

    for idx, (header_text, width) in enumerate(zip(headers, widths)):
        hdr_cells[idx].width = width
        p = hdr_cells[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in (0, 3) else WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(header_text)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        # Estilo de fundo do cabeçalho
        tcPr = hdr_cells[idx]._element.get_or_add_tcPr()
        shd = parse_xml(r'<w:shd {} w:fill="1B365D"/>'.format(nsdecls("w")))
        tcPr.append(shd)

    # Preenchimento das Linhas
    for idx, record in enumerate(drawn_records, start=1):
        row_cells = table.add_row().cells
        values = [
            f"{idx}º",
            record.get("aluno") or record.get("Aluno") or "-",
            record.get("email") or record.get("E-mail") or record.get("nome") or "-",
            record.get("codigo") or record.get("Códigos") or record.get("Codigo") or "-",
            "",  # Coluna em branco para preenchimento manual do prêmio
        ]
        for col_idx, (val, width) in enumerate(zip(values, widths)):
            row_cells[col_idx].width = width
            row_cells[col_idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = row_cells[col_idx].paragraphs[0]
            p.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if col_idx in (0, 3) else WD_ALIGN_PARAGRAPH.LEFT
            )
            run = p.add_run(val)
            run.font.name = "Arial"
            run.font.size = Pt(10)

    doc.add_paragraph()
    doc.add_paragraph()

    # Bloco de Assinaturas
    sig_p = doc.add_paragraph()
    sig_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sig_text = sig_p.add_run(
        "_________________________________________\nComissão Organizadora do Sorteio"
    )
    run_sig_text.font.name = "Arial"
    run_sig_text.font.size = Pt(11)

    doc.save(str(path))
    return path
