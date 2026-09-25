"""Leitura de planilhas (xlsx/csv) e escrita do CSV de saída."""

import csv
import json
import unicodedata
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from .core import MAX_CODE, MIN_CODE
from .runtime import app_root

OUTPUT_FIELDS = ["Aluno", "Nome", "E-mail", "Códigos"]
REPORT_FIELDS = ["Aluno", "Nome", "E-mail", "Códigos", "Status", "Data_Hora", "Detalhes"]
REPORT_PATTERN = "relatorio_envio_*.csv"
UNIFIED_REPORT_PATTERN = "relatorio_envio_unificado_*.csv"

KEY_CODIGOS_EMITIDOS = "codigos_emitidos"
KEY_ALUNOS_RASTREADOS = "alunos_rastreados"
KEY_CODIGOS_SORTEADOS = "codigos_sorteados"
KEY_ALUNOS_SORTEADOS = "alunos_sorteados"
KEY_REPORTS_DIR = "reports_dir"


def _normalize(value: str) -> str:
    text = unicodedata.normalize("NFD", str(value).strip().lower())
    return "".join(char for char in text if unicodedata.category(char) != "Mn")


def normalize_name(value: str) -> str:
    """Normaliza um nome para comparação resiliente entre execuções."""
    return _normalize(value)


def _map_columns(headers: list) -> dict[str, int | None]:
    columns: dict[str, int | None] = {
        "aluno": None,
        "nome": None,
        "email": None,
        "quantidade": None,
    }
    for index, header in enumerate(headers):
        key = _normalize(header)
        if key in {"aluno", "aluno_nome", "student", "nome completo"}:
            columns["aluno"] = index
        elif key in {"nome", "name", "participante", "responsavel financeiro"}:
            columns["nome"] = index
        elif key in {"email", "e-mail", "correio", "e-mail do responsavel financeiro"}:
            columns["email"] = index
        elif key in {"quantidade", "qtd", "qty", "numero", "n"}:
            columns["quantidade"] = index
    return columns


def _detect_delimiter(sample: str) -> str:
    return ";" if sample.count(";") > sample.count(",") else ","


def _read_csv_rows(path: Path) -> list[list[str]]:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            with path.open("r", encoding=encoding, newline="") as fh:
                sample = fh.read(4096)
                fh.seek(0)
                delimiter = _detect_delimiter(sample)
                return list(csv.reader(fh, delimiter=delimiter))
        except UnicodeDecodeError:
            continue
    raise ValueError(f"Não foi possível ler o CSV {path.name}: encoding não reconhecido.")


def _read_xlsx_rows(path: Path) -> list[list]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        if sheet is None:
            raise ValueError(f"A planilha {path.name} não possui planilhas ativas.")
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def _parse_quantity(value) -> int | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        qty = int(value)
        if float(qty) == value and qty >= 1:
            return qty
        return None
    try:
        qty = int(str(value).strip())
    except ValueError:
        return None
    return qty if qty >= 1 else None


def is_valid_email(email: str) -> bool:
    return "@" in email and "." in email.split("@")[-1]


def read_spreadsheet(path: Path) -> tuple[list[dict], list[str]]:
    """Lê a planilha e devolve registros válidos e avisos das linhas ignoradas."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        rows = _read_xlsx_rows(path)
    elif suffix == ".csv":
        rows = _read_csv_rows(path)
    else:
        raise ValueError(f"Formato não suportado: {path.name}")

    if not rows:
        raise ValueError("A planilha está vazia.")

    indexes = _map_columns(rows[0])
    missing = [label for label, index in indexes.items() if index is None]
    if missing:
        raise ValueError(f"Colunas obrigatórias não encontradas: {', '.join(missing)}.")
    columns = {label: index for label, index in indexes.items() if index is not None}

    records = []
    warnings = []
    for line_no, row in enumerate(rows[1:], start=2):
        if not any(str(cell).strip() for cell in row):
            continue
        aluno = str(row[columns["aluno"]]).strip()
        name = str(row[columns["nome"]]).strip()
        email = str(row[columns["email"]]).strip()
        quantity_text = row[columns["quantidade"]]
        quantity = _parse_quantity(quantity_text)
        if not aluno:
            warnings.append(f"Linha {line_no}: aluno vazio, ignorada.")
            continue
        if not name:
            warnings.append(f"Linha {line_no}: nome vazio, ignorada.")
            continue
        if quantity is None:
            warnings.append(f"Linha {line_no}: quantidade inválida ({quantity_text}), ignorada.")
            continue
        email_valid = is_valid_email(email)
        if not email_valid:
            warnings.append(
                f"Linha {line_no}: e-mail inválido ({email or 'vazio'}), mantido para revisão."
            )
        records.append(
            {
                "aluno": aluno,
                "nome": name,
                "email": email,
                "quantidade": quantity,
                "email_valido": email_valid,
            }
        )
    return records, warnings


def filter_new_students(
    records: list[dict], tracked_students: set[str]
) -> tuple[list[dict], list[str]]:
    """Devolve os registros cujos alunos ainda não foram rastreados e os ignorados."""
    fresh = []
    skipped = []
    for record in records:
        key = _normalize(str(record.get("aluno", "")))
        if key and key in tracked_students:
            skipped.append(record.get("aluno", ""))
        else:
            fresh.append(record)
    return fresh, skipped


def write_output_csv(path: Path, rows: list[dict]) -> None:
    """Grava o CSV de saída com BOM e delimitador ';' para o Excel pt-BR."""
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=OUTPUT_FIELDS, delimiter=";")
        writer.writeheader()
        for row in rows:
            writer.writerow({field: row.get(field, "") for field in OUTPUT_FIELDS})


def default_output_path(input_path: Path) -> Path:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return Path(input_path).parent / f"codigos_sorteados_{timestamp}.csv"


def write_report_csv(path: Path, rows: list[dict]) -> None:
    """Grava o CSV de relatório de envio com BOM e delimitador ';'."""
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=REPORT_FIELDS, delimiter=";")
        writer.writeheader()
        writer.writerows(rows)


def default_report_path(input_path: Path) -> Path:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return default_reports_dir() / f"relatorio_envio_{timestamp}.csv"


def default_reports_dir() -> Path:
    """Pasta configurada para relatórios ou o primeiro fallback existente.

    Lê a chave ``reports_dir`` do ``config.json`` da raiz. Quando ela existe
    e aponta para uma pasta válida, devolve esse caminho. Caso contrário,
    devolve a primeira pasta existente entre ``cache/``, ``enviados/`` e a
    raiz do aplicativo (nesta ordem), o que nem sempre precisa existir.
    """
    configurada = configured_reports_dir()
    if configurada is not None:
        return configurada
    raiz = app_root()
    for candidato in (raiz / "cache", raiz / "enviados", raiz):
        if candidato.exists():
            return candidato
    return raiz


def configured_reports_dir() -> Path | None:
    """Devolve a pasta de relatórios salva no ``config.json``, se existir."""
    config = default_config_path()
    if not config.exists():
        return None
    try:
        data = json.loads(config.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return None
    if not isinstance(data, dict):
        return None
    valor = data.get(KEY_REPORTS_DIR)
    if isinstance(valor, str) and valor and Path(valor).exists():
        return Path(valor)
    return None


def set_reports_dir(path: Path | str) -> Path:
    """Cria a pasta de relatórios e a persiste no ``config.json``.

    Devolve o caminho normalizado. Mantém as demais seções do config.json.
    """
    destino = Path(path)
    destino.mkdir(parents=True, exist_ok=True)
    config = default_config_path()
    data: dict = {}
    if config.exists():
        try:
            data = json.loads(config.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, OSError):
            data = {}
    if not isinstance(data, dict):
        data = {}
    data[KEY_REPORTS_DIR] = str(destino)
    config.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    return destino


def latest_unified_report_path(base_dir: Path | None = None) -> Path | None:
    """Devolve o relatório unificado mais recente nas pastas de relatório.

    Com ``base_dir`` informado, procura por ``cache/`` e pela pasta base,
    devolvendo o mais recente entre as duas. No padrão, procura primeiro na
    pasta de relatórios configurada e, se ela estiver vazia, nas pastas
    ``cache/``, ``enviados/`` e raiz do aplicativo — sempre o mais recente.
    Devolve ``None`` se não houver nenhum.
    """
    if base_dir is not None:
        raiz = Path(base_dir)
        diretorios: list[Path] = [raiz / "cache", raiz]
        candidatos = _unified_candidates(diretorios)
        candidatos.sort(key=lambda arquivo: arquivo.stat().st_mtime, reverse=True)
        return candidatos[0] if candidatos else None
    raiz = app_root()
    configurada = configured_reports_dir()
    alvos = [configurada] if configurada else []
    alvos += [raiz / "cache", raiz / "enviados", raiz]
    for diretorio in alvos:
        if not diretorio.exists():
            continue
        candidatos = _unified_candidates([diretorio])
        if candidatos:
            candidatos.sort(key=lambda arquivo: arquivo.stat().st_mtime, reverse=True)
            return candidatos[0]
    return None


def _unified_candidates(diretorios: list[Path]) -> list[Path]:
    candidatos: list[Path] = []
    vistos: set[Path] = set()
    for diretorio in diretorios:
        if not diretorio.exists() or diretorio in vistos:
            continue
        vistos.add(diretorio)
        candidatos.extend(diretorio.glob(UNIFIED_REPORT_PATTERN))
    return candidatos


def unify_reports(report_dir: Path | None = None) -> Path | None:
    """Unifica todos os relatórios de envio em um ``relatorio_envio_unificado_*.csv``.

    Percorre os ``relatorio_envio_*.csv`` do diretório informado (padrão
    a pasta de relatórios configurada), sem incluir os já unificados, e grava
    o novo unificado com a soma de todas as linhas. Devolve o caminho gerado
    ou ``None`` se não houver relatórios com registros válidos.
    """
    diretorio = Path(report_dir) if report_dir else default_reports_dir()
    if not diretorio.exists():
        return None
    arquivos = sorted(
        arquivo
        for arquivo in diretorio.glob(REPORT_PATTERN)
        if not arquivo.name.startswith("relatorio_envio_unificado_")
    )
    linhas: list[dict] = []
    for arquivo in arquivos:
        try:
            linhas.extend(_read_report_rows(arquivo))
        except (OSError, ValueError):
            continue
    if not linhas:
        return None
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    saida = diretorio / f"relatorio_envio_unificado_{timestamp}.csv"
    write_report_csv(saida, linhas)
    return saida


def _read_report_rows(path: Path) -> list[dict]:
    """Lê as linhas não vazias de um relatório de envio (utf-8-sig, ';')."""
    linhas: list[dict] = []
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh, delimiter=";")
        for linha in reader:
            valores = [(linha.get(campo) or "").strip() for campo in REPORT_FIELDS]
            if any(valores):
                linhas.append(dict(zip(REPORT_FIELDS, valores)))
    return linhas


def read_unified_pool(path: Path) -> tuple[set[str], set[int]]:
    """Lê o unificado e devolve (nomes de alunos, códigos) das linhas com Status=Sucesso.

    Cada linha agrupa um aluno com seus códigos; o conjunto de códigos é a
    união de todos os códigos enviados com sucesso.
    """
    nomes: set[str] = set()
    codigos: set[int] = set()
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh, delimiter=";")
        for linha in reader:
            if (linha.get("Status") or "").strip() != "Sucesso":
                continue
            aluno = (linha.get("Aluno") or "").strip()
            if aluno:
                nomes.add(aluno)
            for codigo in (linha.get("Códigos") or "").split(","):
                codigo = codigo.strip()
                if not codigo:
                    continue
                try:
                    numero = int(codigo)
                except ValueError:
                    continue
                if MIN_CODE <= numero <= MAX_CODE:
                    codigos.add(numero)
    return nomes, codigos


def read_unified_pairs(path: Path) -> dict[int, str]:
    """Lê o unificado e devolve {código: aluno} das linhas com Status=Sucesso.

    Cada linha vincula um aluno aos seus códigos; o código é mapeado para o
    aluno que o recebeu. Em caso de duplicidade, o primeiro vínculo vence.
    """
    pares: dict[int, str] = {}
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh, delimiter=";")
        for linha in reader:
            if (linha.get("Status") or "").strip() != "Sucesso":
                continue
            aluno = (linha.get("Aluno") or "").strip()
            if not aluno:
                continue
            for codigo in (linha.get("Códigos") or "").split(","):
                codigo = codigo.strip()
                if not codigo:
                    continue
                try:
                    numero = int(codigo)
                except ValueError:
                    continue
                if MIN_CODE <= numero <= MAX_CODE and numero not in pares:
                    pares[numero] = aluno
    return pares


def default_config_path() -> Path:
    return app_root() / "config.json"


def migrate_legacy_config(config_path: Path) -> bool:
    """Migra os registros JSON antigos para o config.json.

    Devolve True quando a migração aconteceu. Os arquivos legados
    (codigos_emitidos.json, alunos_rastreados.json, codigos_sorteados.json e
    alunos_sorteados.json) são lidos e removidos. Quando algum não existe, o
    valor correspondente vira lista vazia. A migração é ignorada caso o
    config.json já exista, preservando edições posteriores.
    """
    if Path(config_path).exists():
        return False
    base = app_root()
    legacy = {
        KEY_CODIGOS_EMITIDOS: ("codigos_emitidos.json", int),
        KEY_ALUNOS_RASTREADOS: ("alunos_rastreados.json", str),
        KEY_CODIGOS_SORTEADOS: ("codigos_sorteados.json", int),
        KEY_ALUNOS_SORTEADOS: ("alunos_sorteados.json", str),
    }
    data: dict[str, list] = {}
    for section, (filename, cast) in legacy.items():
        path = base / filename
        values: list = []
        if path.exists():
            try:
                values = json.loads(path.read_text(encoding="utf-8"))
            except (json.JSONDecodeError, OSError):
                values = []
        data[section] = sorted({cast(value) for value in values})
    Path(config_path).write_text(
        json.dumps(data, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    for _, (filename, _) in legacy.items():
        path = base / filename
        if path.exists():
            path.unlink()
    return True


def excluded_codes(
    pares: dict[int, str],
    drawn_codes: set[int],
    drawn_students: set[str],
) -> set[int]:
    """Códigos a excluir: os sorteados diretamente mais os de alunos sorteados."""
    excluded = set(drawn_codes)
    normalized = {normalize_name(name) for name in drawn_students}
    for code, aluno in pares.items():
        if normalize_name(aluno) in normalized:
            excluded.add(code)
    return excluded
