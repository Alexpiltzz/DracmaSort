"""Testes unitários para o gerador de documentos Word (.docx)."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from code_gen.docx_export import generate_draw_report_docx


def test_generate_draw_report_docx_cria_arquivo(tmp_path: Path):
    saida = tmp_path / "teste_ata.docx"
    registros = [
        {"aluno": "Maria Silva", "email": "maria@exemplo.com", "codigo": "0042"},
        {"aluno": "João Souza", "email": "joao@exemplo.com", "codigo": "0100"},
    ]

    resultado = generate_draw_report_docx(saida, registros)

    assert resultado.exists()
    assert resultado == saida

    # Ler o documento gerado com a biblioteca python-docx
    doc = Document(resultado)
    assert len(doc.tables) == 1

    tabela = doc.tables[0]
    # 1 linha de cabeçalho + 2 registros
    assert len(tabela.rows) == 3

    cabecalho = [cell.text for cell in tabela.rows[0].cells]
    assert cabecalho == ["№", "Aluno", "E-mail / Responsável", "Código Sorteado", "Prêmio Recebido"]

    primeiro_registro = [cell.text for cell in tabela.rows[1].cells]
    assert primeiro_registro[0] == "1º"
    assert primeiro_registro[1] == "Maria Silva"
    assert primeiro_registro[2] == "maria@exemplo.com"
    assert primeiro_registro[3] == "0042"
    assert primeiro_registro[4] == ""  # Prêmio Recebido em branco para preenchimento manual


def test_generate_draw_report_docx_sem_registros(tmp_path: Path):
    saida = tmp_path / "ata_vazia.docx"
    resultado = generate_draw_report_docx(saida, [])

    assert resultado.exists()
    doc = Document(resultado)
    assert len(doc.tables) == 1
    assert len(doc.tables[0].rows) == 1  # apenas cabeçalho
