"""Ferramentas de desenvolvimento — terminal interativo.

Execução:
    uv run python main_dev.py
"""

from __future__ import annotations

import csv
import unicodedata
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from code_gen.io import REPORT_FIELDS, is_valid_email, write_report_csv
from code_gen.runtime import app_root

CACHE_DIR = app_root() / "cache"
REPORT_PATTERN = "relatorio_envio_*.csv"
UNIFICADO_PATTERN = "relatorio_envio_unificado_*.csv"
ALUNOS_PLANILHA_PATTERN = "alunos_planilha_*.csv"
STATUS_ORDEM = ["Sucesso", "Falha", "Rejeitado", "Simulado"]

COLUNAS_ALVO = {
    "aluno": "nomecompleto",
    "responsavel": "responsavelfinanceiro",
    "email": "emaildoresponsavelfinanceiro",
}
FIELDS_ALUNOS = ["Aluno", "Nome", "E-mail"]
COMPARATIVO_FIELDS = [
    "Aluno",
    "Responsavel_Relatorio",
    "Responsavel_Cadastrado",
    "Situacao",
    "Status_Envio",
]
ORDEM_SITUACAO = {
    "Responsável divergente": 0,
    "Sem correspondência no sistema": 1,
    "Não consta no relatório": 2,
    "Igual": 3,
}


def _ler_report(path: Path) -> list[dict]:
    """Lê um relatório de envio e devolve as linhas limpas e normalizadas."""
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh, delimiter=";")
        linhas = []
        for linha in reader:
            if any((valor or "").strip() for valor in linha.values()):
                linhas.append({campo: (linha.get(campo) or "").strip() for campo in REPORT_FIELDS})
    return linhas


def _contar_codigos(linha: dict) -> int:
    """Conta a quantidade de códigos presentes no campo 'Códigos'."""
    codigos = linha.get("Códigos", "") or ""
    return sum(1 for codigo in codigos.split(",") if codigo.strip())


def unificar_relatorios(cache_dir: Path = CACHE_DIR) -> Path | None:
    """Unifica todos os relatórios de envio do cache em um único CSV e exibe o resumo."""
    diretorio = Path(cache_dir)
    arquivos = sorted(diretorio.glob(REPORT_PATTERN))
    if not arquivos:
        print(f"Nenhum relatório encontrado em {diretorio} (padrão {REPORT_PATTERN!r}).")
        return None

    linhas: list[dict] = list()
    falhados = []
    for arquivo in arquivos:
        try:
            linhas.extend(_ler_report(arquivo))
        except Exception as exc:  # noqa: BLE001 — ferramenta de dev, mensagem descritiva basta
            print(f"  [aviso] não foi possível ler {arquivo.name}: {exc}")
            falhados.append(arquivo.name)

    if not linhas:
        print("Nenhum registro válido encontrado nos relatórios.")
        return None

    total = len(linhas)
    status_counts = Counter(linha.get("Status", "") or "(sem status)" for linha in linhas)
    total_codigos = sum(_contar_codigos(linha) for linha in linhas)
    emails_unicos = {
        linha.get("E-mail", "").strip() for linha in linhas if linha.get("E-mail", "").strip()
    }
    falhas = Counter(
        (linha.get("Status", ""), (linha.get("Detalhes", "") or "")[:140])
        for linha in linhas
        if linha.get("Status") != "Sucesso"
    )

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    saida = diretorio / f"relatorio_envio_unificado_{timestamp}.csv"
    write_report_csv(saida, linhas)

    separador()
    print("=== Resumo — relatórios de envio unificados ===")
    print(
        f"Arquivos lidos:                {len(arquivos) - len(falhados)}/{len(arquivos)}"
        + (f"  (falhas de leitura: {', '.join(falhados)})" if falhados else "")
    )
    print(f"Total de registros:            {total}")
    print(f"E-mails distintos:             {len(emails_unicos)}")
    print(f"Códigos emitidos no total:     {total_codigos}")

    print("\nPor status:")
    for chave in STATUS_ORDEM:
        _linha_status(chave, status_counts.get(chave, 0), total)
    outras = sorted({chave for chave in status_counts if chave not in STATUS_ORDEM})
    for chave in outras:
        _linha_status(chave, status_counts[chave], total)

    print("\nMotivos mais comuns fora de sucesso:")
    if falhas:
        for (status, detalhe), qtde in falhas.most_common(5):
            print(f"  [{status}] {detalhe!r} — {qtde}x")
    else:
        print("  nenhum.")

    print(f"\nArquivo unificado gerado: {saida}")
    return saida


def _linha_status(chave: str, qtde: int, total: int) -> None:
    pct = (qtde / total * 100) if total else 0.0
    print(f"  {chave:<12}{qtde:>7}{pct:>8.1f}%")


def _normalizar_coluna(texto: str) -> str:
    """Normaliza um cabeçalho para comparação (minúsculas, sem acentos/hífens)."""
    sem_acentos = "".join(
        char for char in unicodedata.normalize("NFD", texto) if unicodedata.category(char) != "Mn"
    )
    return sem_acentos.strip().lower().replace("-", "").replace(" ", "")


def _localizar_cabecalho(linhas: list[list]) -> tuple[int | None, dict[str, int] | None]:
    """Encontra a linha de cabeçalho com as colunas de interesse e o mapa coluna -> índice."""
    for pos, linha in enumerate(linhas[:30], start=1):
        mapa = {
            _normalizar_coluna(str(celula)): indice
            for indice, celula in enumerate(linha)
            if celula is not None
        }
        if all(valor in mapa for valor in COLUNAS_ALVO.values()):
            return pos, mapa
    return None, None


def _ler_linhas_planilha(path: Path) -> list[list]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        if ws is None:
            raise ValueError(f"A planilha {path.name} não possui planilhas ativas.")
        return [list(linha) for linha in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def _detectar_planilhas(cache_dir: Path) -> list[Path]:
    """Lista os .xlsx do cache cujo cabeçalho contém as colunas de interesse."""
    detectadas = []
    for arquivo in sorted(cache_dir.glob("*.xlsx")):
        linhas = _ler_linhas_planilha(arquivo)[:30]
        _, mapa = _localizar_cabecalho(linhas)
        if mapa is not None:
            detectadas.append(arquivo)
    return detectadas


def _escolher_arquivo(candidatas: list[Path]) -> Path | None:
    if len(candidatas) == 1:
        return candidatas[0]
    print("Planilhas com as colunas de interesse encontradas no cache:")
    for indice, arquivo in enumerate(candidatas, start=1):
        print(f"  {indice}) {arquivo.name}")
    while True:
        try:
            escolha = input(f"Escolha 1-{len(candidatas)} (0 para cancelar): ").strip()
        except (EOFError, KeyboardInterrupt):
            return None
        if escolha == "0":
            return None
        if escolha.isdigit() and 1 <= int(escolha) <= len(candidatas):
            return candidatas[int(escolha) - 1]
        print("Opção inválida.")


def _ler_alunos_da_planilha(path: Path) -> tuple[list[dict], list[str]]:
    """Extrai Aluno/Nome/E-mail da planilha base e devolve registros e avisos."""
    linhas = _ler_linhas_planilha(path)

    pos_cabecalho, mapa = _localizar_cabecalho(linhas)
    if pos_cabecalho is None or mapa is None:
        raise ValueError(f"Cabeçalho não encontrado em {path.name}.")

    indices = {chave: mapa[valor] for chave, valor in COLUNAS_ALVO.items()}
    registros = []
    avisos = []
    for numero, linha in enumerate(linhas[pos_cabecalho:], start=pos_cabecalho + 1):
        if not any(celula is not None and str(celula).strip() for celula in linha):
            continue

        def celula(chave: str) -> str:
            indice = indices[chave]
            if indice >= len(linha) or linha[indice] is None:
                return ""
            return str(linha[indice]).strip()

        aluno, responsavel, email = celula("aluno"), celula("responsavel"), celula("email")
        if not (aluno or responsavel or email):
            continue
        if not aluno:
            avisos.append(f"Linha {numero}: aluno sem nome, mantido.")
        if not responsavel:
            avisos.append(f"Linha {numero}: responsável financeiro vazio para {aluno or email}.")
        if not email:
            avisos.append(f"Linha {numero}: e-mail vazio para {aluno or responsavel}.")
        elif not is_valid_email(email):
            avisos.append(f"Linha {numero}: e-mail aparentemente inválido ({email}).")
        registros.append({"Aluno": aluno, "Nome": responsavel, "E-mail": email})
    return registros, avisos


def extrair_alunos(cache_dir: Path = CACHE_DIR) -> Path | None:
    """Extrai as colunas Aluno/Nome/E-mail da planilha base para um CSV no cache."""
    diretorio = Path(cache_dir)
    candidatas = _detectar_planilhas(diretorio)
    if not candidatas:
        print(
            "Nenhuma planilha no cache possui as colunas "
            "'Nome Completo', 'Responsável Financeiro' e 'E-Mail do Responsável Financeiro'."
        )
        return None

    arquivo = _escolher_arquivo(candidatas)
    if arquivo is None:
        return None

    registros, avisos = _ler_alunos_da_planilha(arquivo)
    if not registros:
        print(f"Nenhum registro extraído de {arquivo.name}.")
        return None

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    saida = diretorio / f"alunos_planilha_{timestamp}.csv"
    with saida.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=FIELDS_ALUNOS, delimiter=";")
        writer.writeheader()
        writer.writerows(registros)

    emails = [registro["E-mail"] for registro in registros if registro["E-mail"]]
    separador()
    print("=== Resumo — extração da planilha base ===")
    print(f"Planilha:                      {arquivo.name}")
    print(f"Total de registros:            {len(registros)}")
    print(f"E-mails preenchidos:           {len(emails)}")
    print(f"E-mails distintos:             {len(set(emails))}")
    print(
        f"E-mails aparentemente inválidos: "
        f"{sum(1 for email in emails if not is_valid_email(email))}"
    )
    print(
        f"Responsáveis sem nome:          "
        f"{sum(1 for registro in registros if not registro['Nome'])}"
    )
    print(
        f"Registros sem nome do aluno:   "
        f"{sum(1 for registro in registros if not registro['Aluno'])}"
    )
    if avisos:
        print(f"\nAvisos ({len(avisos)}):")
        for aviso in avisos[:10]:
            print(f"  {aviso}")
        if len(avisos) > 10:
            print(f"  ... e mais {len(avisos) - 10}.")

    print(f"\nArquivo extraído gerado: {saida}")
    return saida


def separador(caractere: str = "—") -> None:
    print(caractere * 56)


def _escolher_arquivo_por_padrao(diretorio: Path, padrao: str, descricao: str) -> Path | None:
    """Escolhe um arquivo do cache por padrão de nome (mais recente primeiro)."""
    arquivos = sorted(
        diretorio.glob(padrao), key=lambda arquivo: arquivo.stat().st_mtime, reverse=True
    )
    if not arquivos:
        print(f"Nenhum {descricao} encontrado em {diretorio} (padrão {padrao!r}).")
        return None
    if len(arquivos) == 1:
        print(f"Usando {descricao}: {arquivos[0].name}")
        return arquivos[0]
    print(f"{descricao.capitalize()}s encontrados:")
    for indice, arquivo in enumerate(arquivos, start=1):
        print(f"  {indice}) {arquivo.name}")
    while True:
        try:
            escolha = input(f"Escolha 1-{len(arquivos)} (0 para cancelar): ").strip()
        except (EOFError, KeyboardInterrupt):
            return None
        if escolha == "0":
            return None
        if escolha.isdigit() and 1 <= int(escolha) <= len(arquivos):
            return arquivos[int(escolha) - 1]
        print("Opção inválida.")


def _codigos_de(registro: dict) -> set[str]:
    return {
        codigo.strip()
        for codigo in (registro.get("Códigos", "") or "").split(",")
        if codigo.strip()
    }


def analisar_alunos_repetidos(cache_dir: Path = CACHE_DIR) -> None:
    """Analisa alunos repetidos no unificado e compara os códigos enviados em cada registro."""
    diretorio = Path(cache_dir)
    arquivo = _escolher_arquivo_por_padrao(diretorio, UNIFICADO_PATTERN, "relatório unificado")
    if arquivo is None:
        return

    linhas = _ler_report(arquivo)
    if not linhas:
        print("Nenhum registro no relatório unificado.")
        return

    por_aluno: dict[str, list[dict]] = {}
    for linha in linhas:
        por_aluno.setdefault(linha.get("Aluno", ""), []).append(linha)

    repetidos = {aluno: registros for aluno, registros in por_aluno.items() if len(registros) > 1}
    iguais: list[tuple[str, list[dict], set[str]]] = []
    diferentes: list[tuple[str, list[dict], list[set[str]]]] = []
    for aluno, registros in sorted(repetidos.items()):
        combinacoes = [_codigos_de(registro) for registro in registros]
        if all(combinacao == combinacoes[0] for combinacao in combinacoes[1:]):
            iguais.append((aluno, registros, combinacoes[0]))
        else:
            diferentes.append((aluno, registros, combinacoes))

    separador()
    print("=== Resumo — alunos repetidos ===")
    print(f"Relatório analisado:            {arquivo.name}")
    print(f"Total de registros:             {len(linhas)}")
    print(f"Alunos distintos:               {len(por_aluno)}")
    print(f"Alunos com mais de um registro: {len(repetidos)}")
    print(f"  > com os mesmos códigos:      {len(iguais)}")
    print(f"  > com códigos diferentes:     {len(diferentes)}")

    detalhes = iguais + diferentes
    if detalhes:
        print("\nDetalhes (até 30):")
        for aluno, registros, combinacoes in detalhes[:30]:
            statuses = ", ".join(sorted({registro.get("Status", "?") for registro in registros}))
            if isinstance(combinacoes, set):
                bloco = ", ".join(sorted(combinacoes))
                descricao = f"mesmos códigos [{bloco}]"
            else:
                blocos = " | ".join(", ".join(sorted(c)) for c in combinacoes)
                descricao = f"códigos diferentes [{blocos}]"
            print(f"  {aluno} — {len(registros)}x — {descricao} | status: {statuses}")
        resto = len(detalhes) - 30
        if resto > 0:
            print(f"  ... e mais {resto} alunos repetidos.")


def _ler_planilha_alunos(path: Path) -> list[dict]:
    """Lê um CSV de alunos extraído (Aluno;Nome;E-mail)."""
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh, delimiter=";")
        return [
            {campo: (linha.get(campo) or "").strip() for campo in FIELDS_ALUNOS}
            for linha in reader
            if any((valor or "").strip() for valor in linha.values())
        ]


def comparar_aluno_responsavel(cache_dir: Path = CACHE_DIR) -> None:
    """Compara os pares aluno+responsável entre o unificado e a planilha base."""
    diretorio = Path(cache_dir)
    unificado = _escolher_arquivo_por_padrao(diretorio, UNIFICADO_PATTERN, "relatório unificado")
    if unificado is None:
        return
    planilha = _escolher_arquivo_por_padrao(
        diretorio, ALUNOS_PLANILHA_PATTERN, "planilha de alunos"
    )
    if planilha is None:
        return

    linhas_unificado = _ler_report(unificado)
    linhas_planilha = _ler_planilha_alunos(planilha)
    if not linhas_unificado or not linhas_planilha:
        print("Um dos arquivos está vazio; comparação cancelada.")
        return

    def chave(aluno: str, nome: str) -> tuple[str, str]:
        return _normalizar_coluna(aluno), _normalizar_coluna(nome)

    pares_unificado: dict[tuple[str, str], tuple[str, str]] = {}
    for linha in linhas_unificado:
        pares_unificado.setdefault(
            chave(linha["Aluno"], linha["Nome"]), (linha["Aluno"], linha["Nome"])
        )

    pares_planilha: dict[tuple[str, str], tuple[str, str]] = {}
    responsaveis_por_aluno: dict[str, set[str]] = {}
    for linha in linhas_planilha:
        chave_norm = chave(linha["Aluno"], linha["Nome"])
        pares_planilha.setdefault(chave_norm, (linha["Aluno"], linha["Nome"]))
        responsaveis_por_aluno.setdefault(_normalizar_coluna(linha["Aluno"]), set()).add(
            linha["Nome"]
        )

    conjunto_unificado = set(pares_unificado)
    conjunto_planilha = set(pares_planilha)
    iguais = sorted(conjunto_unificado & conjunto_planilha)

    somente_unificado = sorted(conjunto_unificado - conjunto_planilha)
    somente_planilha = sorted(conjunto_planilha - conjunto_unificado)

    divergentes = []
    for chave_norm in somente_unificado:
        aluno_norm, _ = chave_norm
        if aluno_norm in responsaveis_por_aluno:
            divergentes.append(chave_norm)

    somente_unificado_set = set(somente_unificado)
    sem_sistema = sorted(
        chave_norm
        for chave_norm in somente_unificado_set
        if chave_norm[0] not in responsaveis_por_aluno
    )

    status_por_aluno: dict[str, set[str]] = {}
    for linha in linhas_unificado:
        status_por_aluno.setdefault(_normalizar_coluna(linha["Aluno"]), set()).add(
            linha.get("Status") or "-"
        )

    registro_csv: list[dict] = []
    for chave_norm in conjunto_unificado:
        aluno, nome = pares_unificado[chave_norm]
        aluno_norm = chave_norm[0]
        status = ", ".join(sorted(status_por_aluno.get(aluno_norm, {"-"})))
        if chave_norm in conjunto_planilha:
            situacao, cadastrado = "Igual", nome
        elif aluno_norm in responsaveis_por_aluno:
            situacao = "Responsável divergente"
            cadastrado = ", ".join(sorted(responsaveis_por_aluno[aluno_norm]))
        else:
            situacao, cadastrado = "Sem correspondência no sistema", ""
        registro_csv.append(
            {
                "Aluno": aluno,
                "Responsavel_Relatorio": nome,
                "Responsavel_Cadastrado": cadastrado,
                "Situacao": situacao,
                "Status_Envio": status,
            }
        )

    for chave_norm in somente_planilha:
        aluno, nome = pares_planilha[chave_norm]
        registro_csv.append(
            {
                "Aluno": aluno,
                "Responsavel_Relatorio": "",
                "Responsavel_Cadastrado": nome,
                "Situacao": "Não consta no relatório",
                "Status_Envio": "-",
            }
        )

    registro_csv.sort(
        key=lambda linha: (
            ORDEM_SITUACAO[linha["Situacao"]],
            _normalizar_coluna(linha["Aluno"]),
        )
    )

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    saida = diretorio / f"comparativo_aluno_responsavel_{timestamp}.csv"
    with saida.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=COMPARATIVO_FIELDS, delimiter=";")
        writer.writeheader()
        writer.writerows(registro_csv)

    separador()
    print("=== Resumo — comparativo aluno + responsável ===")
    print(f"Relatório unificado:              {unificado.name}")
    print(f"Planilha base:                    {planilha.name}")
    print(f"Pares no relatório unificado:     {len(conjunto_unificado)}")
    print(f"Pares na planilha base:           {len(conjunto_planilha)}")
    print(f"Pares iguais (mesmo responsável): {len(iguais)}")
    print(f"Responsável divergente:           {len(divergentes)}")
    print(f"Sem correspondência no sistema:   {len(sem_sistema)}")
    print(f"Não constam no relatório:         {len(somente_planilha)}")

    def _lista(titulo: str, chaves: list[tuple[str, str]]) -> None:
        print(f"\n{titulo}:")
        if not chaves:
            print("  nenhum.")
            return
        for chave_norm in chaves[:30]:
            aluno, nome = pares_unificado.get(chave_norm) or pares_planilha.get(
                chave_norm, ("", "")
            )
            print(f"  {aluno} | {nome}")
        resto = len(chaves) - 30
        if resto > 0:
            print(f"  ... e mais {resto}.")

    print("\nResponsável divergente (relatório x cadastrado):")
    if divergentes:
        for chave_norm in divergentes[:30]:
            aluno, nome_relatorio = pares_unificado[chave_norm]
            cadastrado = ", ".join(sorted(responsaveis_por_aluno[chave_norm[0]]))
            print(f"  {aluno} — relatório: {nome_relatorio} | cadastrado: {cadastrado}")
        resto = len(divergentes) - 30
        if resto > 0:
            print(f"  ... e mais {resto}.")
    else:
        print("  nenhum.")

    _lista("Alunos do relatório sem correspondência no sistema", sem_sistema)
    _lista("Alunos do sistema que não constam no relatório", somente_planilha)

    print(f"\nCSV completo gerado: {saida}")


def _mostrar_menu() -> None:
    print("\n1) Unificar relatórios de envio")
    print("2) Extrair alunos da planilha base")
    print("3) Analisar alunos repetidos (unificado)")
    print("4) Comparar aluno + responsável (unificado vs planilha)")
    print("0) Sair")


def main() -> int:
    print("=== Ferramentas de Desenvolvimento — Giveaways_Tools ===")
    while True:
        _mostrar_menu()
        try:
            opcao = input("Opção: ").strip().lower()
        except (EOFError, KeyboardInterrupt):
            print("\nAté logo!")
            return 0
        if opcao == "1":
            separador()
            unificar_relatorios()
        elif opcao == "2":
            separador()
            extrair_alunos()
        elif opcao == "3":
            separador()
            analisar_alunos_repetidos()
        elif opcao == "4":
            separador()
            comparar_aluno_responsavel()
        elif opcao in {"0", "sair", "q"}:
            print("Até logo!")
            return 0
        else:
            print("Opção inválida. Tente novamente.")


if __name__ == "__main__":
    raise SystemExit(main())
